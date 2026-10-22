"""Paper-first exploratory admission using explicit native-solver assumptions.

Preserves archived workflows and sources. Fixed treatments are tried in declared
order using calibration-only feasibility, never assessment scores. The numerical
equations, native steps and Hazen percentile boxes remain unchanged.
"""
from dataclasses import replace
from pathlib import Path
import re
import numpy as np
import pandas as pd
from openpyxl import load_workbook
from workflows import reservoir_v2 as w
from strategies.representative_years import select_representative_years, fit_selected_years, RepresentativeYearObjective
from strategies.fleet_comparison_v2 import comparison_selections
from dztr.calibration.sampled import _observations
from dztr.calibration.candidates import parameters_to_vector
from dztr.integrations.world_reservoirs import _forcing
from dztr.schemas import ReservoirParameters
from dztr.simulation.sampled import SamplingSettings
from tools.fleet_admission import ordered_candidate

ROOT = w.ROOT
TABLE = ROOT / 'data/paper_cohort_v1/hess-23-3735-2019-t01.xlsx'
PAPER_IDS = {
 'American Falls':'ex13_American_falls','Andijan1':'ex34_Andijan','Bhumibol':'ex17_Bhumibol',
 'Big Horn':'ex30_Big_horn','Bull Lake1':'ex35_Bull_Lake','Canyon Ferry1':'ex7_CanyonFerry',
 'Chardara':'ex36_Chardara','Charvak1':'ex37_Charvak','Dickson':'ex28_Dickson',
 'E. B. Campbell1':'ex29_EB_Campbell','Flaming Gorge1':'ex8_Flaming_gorge','Fort Peck1':'ex3_Fort_peck',
 'Fort Randall':'ex4_Fort_randall','Gardiner':'ex14_Gardiner','Garrison1':'ex5_Garrison',
 'Ghost':'ex26_Ghost','Glen Canyon1':'ex1_Glen_canyon','Grand Coulee':'ex2_Grand_coulee',
 'High Aswan':'ex18_High_aswan','Amistad (Int. Amistad)':'ex9_Amistad',
 'Falcon International':'ex10_Falcon','Kayrakkum1':'ex38_Kayrakkum','Navajo':'ex33_Navajo',
 'Nurek':'ex39_Nurek','Oahe Dam1':'ex6_Oahe','Oldman River':'ex24_Oldman',
 'Oroville1':'ex32_Oroville','Palisades':'ex11_Palisades','Seminoe':'ex40_Semione',
 'Sirikit':'ex19_Sirikit','St. Mary':'ex20_St_marry','Toktogul1':'ex41_Toktogul',
 'Trinity':'ex12_Trinity','Tyuyamuyun':'ex42_Tyuyamuyun','W. A. C. Bennett':'ex15_Bennett',
 'Waterton':'ex23_Waterton','Yellowtail':'ex16_Yellowtail'}
NATIVE_PRIORITY = ['ex30_Big_horn','ex41_Toktogul','ex36_Chardara','ex38_Kayrakkum',
                   'ex11_Palisades','ex3_Fort_peck','ex10_Falcon','ex13_American_falls']
SOURCE_BLOCKS = {
 'ex7_CanyonFerry':'Confirmed calendar/rate errors and missing reconstructed months: source repair required.',
 'ex12_Trinity':'Known conversion/version and storage-timing repairs: source profile required.',
 'ex26_Ghost':'Unresolved storage coordinate, clock and boundary definitions: source profile required.',
 'ex15_Bennett':'Gross-versus-active storage coordinate remains unresolved; paper capacity is not substituted into that series.',
 'ex8_Flaming_gorge':'Signed inflow requires an explicit loss/net-flux adapter; no clipping or absolute values.',
 'ex33_Navajo':'Signed inflow requires an explicit loss/net-flux adapter; no clipping or absolute values.',
 'ex40_Semione':'Signed inflow requires an explicit loss/net-flux adapter; no clipping or absolute values.'}

def paper_cohort():
    proof = w.read(TABLE.parent / 'download.json')
    if w.sha(TABLE) != proof['sha256']:
        raise ValueError('Published Table 1 changed')
    book = load_workbook(TABLE, read_only=True, data_only=True)
    records = []
    for row_number, values in enumerate(book.worksheets[0].iter_rows(values_only=True), 1):
        if row_number <= 5 or values[1] is None:
            continue  # Header or wrapped Falcon name continuation, never a second reservoir.
        name = re.sub(r'\s+', ' ', values[0]).strip()
        if name not in PAPER_IDS:
            raise ValueError(f'Unmapped published reservoir: {name}')
        capacity = float(re.sub(r'\s+', '', str(values[7]))) * 1e6
        records.append({'paper_order': len(records)+1, 'table_row': row_number,
                        'paper_name': name, 'reservoir_id': PAPER_IDS[name], 'country': values[1],
                        'paper_capacity_m3': capacity, 'paper_longitude_verbatim': values[4],
                        'paper_latitude_verbatim': values[5], 'paper_period_verbatim': values[9]})
    book.close()
    if len(records) != 37 or len({r['reservoir_id'] for r in records}) != 37:
        raise ValueError('Published cohort must have 37 unique reservoirs')
    return records

def floor_candidates(original):
    """Retain the original first, then fixed lower hypotheses; no fitting to scores."""
    return [original] + [v for v in (.1, .05, 0.) if v < original]

def choose_treatment(observations, config, plan, paper_record):
    """Find a fixed calibration-feasible witness; assessment responses are unused."""
    if observations.frame.inflow_m3s.lt(0).any():
        raise ValueError('Signed forcing needs a net-flux/loss adapter; values cannot be clipped')
    sampling = SamplingSettings(**config['sampling'])
    attempts = []
    capacities = [('inherited_nominal_capacity', config['model']['capacity_m3'])]
    # Explicit gross-capacity hypotheses are restricted to these two documented cases.
    if paper_record['reservoir_id'] in {'ex3_Fort_peck','ex10_Falcon'}:
        capacities.append(('published_nominal_gross_capacity_hypothesis', paper_record['paper_capacity_m3']))
    for capacity_kind, capacity in capacities:
        for floor in floor_candidates(config['model']['dead_storage_fraction']):
            constants = dict(config['model'], capacity_m3=capacity, dead_storage_fraction=floor)
            description = {'capacity_kind': capacity_kind, 'capacity_m3': capacity,
                           'dead_storage_fraction': floor, 'floor_is_assumed_model_policy': True,
                           'assessment_used': False}
            try:
                template, bounds = fit_selected_years(observations, plan, constants, allow_feasible_initialization=True)
            except ValueError as error:
                attempts.append({**description,'stage':'parameter_preparation','feasible':False,'reason':str(error)})
                continue
            points = [('prepared_initializer', template)]
            lower = bounds.empirical.monthly_bounds.iloc[:, :6].to_numpy()
            upper = bounds.empirical.monthly_bounds.iloc[:, 6:].to_numpy()
            step_days = config['sampling'].get('fixed_step_days') or 1.
            gap = ordered_candidate(lower, upper, capacity, floor, step_days * 86400, close_critical_gap=True)
            if gap['feasible']:
                candidate = ReservoirParameters(pd.DataFrame(gap['monthly_targets_si'], index=template.monthly_targets.index,
                                               columns=template.monthly_targets.columns), **constants)
                points.append(('sufficient_critical_gap_witness', bounds.validate(parameters_to_vector(candidate), template)))
            for initial_kind, initial in points:
                objective = RepresentativeYearObjective(observations, plan, initial, sampling=sampling, bounds=bounds,
                           policy='nearest_previous_storage', atol_storage_m3=config['atol_storage_m3'], cost_mode='pooled_nse')
                evaluation = objective.evaluate(parameters_to_vector(initial))
                attempt = {**description,'stage':'fixed_calibration_preflight','initializer_kind':initial_kind,
                           'feasible':evaluation.feasible,'reason':evaluation.reason,'metrics':evaluation.metrics}
                if evaluation.simulation is not None:
                    attempt['independent_balance_residual_m3'] = w.independent_balance(evaluation.simulation.frame, initial.inflow_correction)
                attempts.append(attempt)
                if evaluation.feasible:
                    if attempt['independent_balance_residual_m3'] > config['atol_storage_m3']:
                        raise ValueError('Independent balance gate failed')
                    return initial, plan, attempts
    return None, plan, attempts

def prepare(output):
    out = w.fresh_output(output)
    cohort = paper_cohort(); by_id = {r['reservoir_id']:r for r in cohort}
    programme = pd.read_csv(ROOT / 'outputs/source_limit_review_v1_20261013/programme_status.csv')
    programme['in_published_paper'] = programme.reservoir_id.isin(by_id)
    programme['paper_priority'] = np.where(programme.in_published_paper, 1, 2)
    programme.loc[programme.included.eq(False), 'paper_priority'] = 3
    programme['paper_order'] = programme.reservoir_id.map({r['reservoir_id']:r['paper_order'] for r in cohort})
    pd.DataFrame(cohort).to_csv(out / 'paper_cohort.csv', index=False)
    active_paper = programme.loc[programme.in_published_paper & programme.included]
    assert len(active_paper) == 37
    w.write(out / 'cohort.json', {'paper_url':'https://hess.copernicus.org/articles/23/3735/2019/',
            'table_url':w.read(TABLE.parent/'download.json')['url'], 'table_sha256':w.sha(TABLE),
            'records':cohort,'paper_count':37,'active_nonpaper_ids':programme.loc[programme.included & ~programme.in_published_paper,'reservoir_id'].tolist(),
            'exclusion_retained':'ex43_TuyenQuang','paper_metadata_does_not_override_dates_or_coordinates':True})
    ready=[];reviews=[]
    for identifier in NATIVE_PRIORITY:
        record=by_id[identifier];folder=out/identifier;folder.mkdir()
        config_path=ROOT/'configs/reservoirs'/f'{identifier}.json';cfg=w.read(config_path)
        observation_path=ROOT/'data/reservoirs'/identifier/'observations_si.csv'
        if w.sha(observation_path)!=cfg['observations_sha256']:
            raise ValueError('Canonical source changed')
        raw=_observations(pd.read_csv(observation_path,index_col='date',parse_dates=True,float_precision='round_trip'))
        adjusted=_forcing(raw,cfg['inflow_preprocessing'])
        _,annual,_=select_representative_years(adjusted,SamplingSettings(**cfg['sampling']))
        plans,_=comparison_selections(adjusted.frame,annual);plan=plans['mixed_pooled']
        initial,plan,attempts=choose_treatment(adjusted,cfg,plan,record)
        review={'reservoir_id':identifier,'in_published_paper':True,'attempts':attempts,
                'calibration_fraction':plan.roles.count('calibration')/len(plan.roles),
                'selection_sha256':plan.signature,'assessment_used_for_admission':False,
                'source_definitions_certified':False,'canonical_inputs_changed':False}
        w.write(folder/'preflight.json',review)
        if initial is None:
            review['status']='no_declared_native_treatment_passed'
        else:
            adjusted.frame.to_csv(folder/'observations_model_input_si.csv',index_label='date',float_format='%.17g')
            initial.monthly_targets.to_csv(folder/'initializer_si.csv',float_format='%.17g')
            model={'physical_constants':{k:getattr(initial,k) for k in cfg['model']},'sampling':cfg['sampling'],
                   'root_policy':'nearest_previous_storage','atol_storage_m3':cfg['atol_storage_m3'],
                   'quantile_method':'hazen','source_preprocessing_already_applied':True}
            w.write(folder/'model.json',model)
            assumption={'kind':'paper_exploratory_treatment_v1','reservoir_id':identifier,'in_published_paper':True,
                'canonical_config_sha256':w.sha(config_path),'canonical_observations_sha256':w.sha(observation_path),
                'published_table_sha256':w.sha(TABLE),'chosen_treatment':attempts[-1],
                'selection_rule':'first physically and calibration-feasible fixed treatment in declared order; no score ranking',
                'treatment_order':'inherited capacity before paper-capacity hypothesis (Fort Peck/Falcon only); original floor before 0.1/0.05/0 where lower; prepared initializer before sufficient-gap witness',
                'observations_outside_nominal_model_limits':'retained as reported and scored; model states/targets remain physically bounded',
                'nominal_limit_does_not_certify_surcharge':True,'operating_intent_recovered':False,
                'root_policy':'nearest_previous_storage','floor_is_assumed_not_measured_dead_storage':True,
                'inflow_preprocessing_applied_once':cfg['inflow_preprocessing'],
                'observations_dropped_or_clipped':False,'bounds_widened':False,'native_timestep_preserved':True,
                'assessment_used_for_admission':False,'GRAND_ID':None,'parameter_status':'exploratory'}
            w.write(folder/'assumptions.json',assumption)
            entry={'reservoir_id':identifier,'grand_id':None,'observations_csv':'observations_model_input_si.csv',
                   'model_config':'model.json','initializer':{'csv':'initializer_si.csv','sha256':w.sha(folder/'initializer_si.csv'),
                   'basis':'calibration_only','selection_sha256':plan.signature,'provenance':str((folder/'assumptions.json').relative_to(ROOT))}}
            config={'schema_version':2,'year_selection':'mixed','scoring':'pooled',
                    'optimization':{'seeds':[42,43,44],'population':32,'evaluation_budget':512,'workers':3},'reservoirs':[entry]}
            w.write(folder/'workflow.json',config)
            prepared=w.prepare(folder/'workflow.json',folder/'calibration')
            ready.append({'reservoir_id':identifier,'config':str((folder/'workflow.json').relative_to(ROOT)),
                          'output':str(prepared.relative_to(ROOT))})
            review['status']='exploratory_native_preflight_passed'
        reviews.append(review)
        print(identifier,review['status'],flush=True)
    for identifier,reason in SOURCE_BLOCKS.items():
        reviews.append({'reservoir_id':identifier,'in_published_paper':True,'status':'source_or_adapter_required','reason':reason,'new_optimization_evaluations':0})
    lookup={r['reservoir_id']:r['status'] for r in reviews}
    programme['exploratory_treatment_status']=programme.reservoir_id.map(lookup).fillna('previous_status_retained')
    programme.sort_values(['paper_priority','paper_order','reservoir_id']).to_csv(out/'prioritized_programme.csv',index=False)
    w.write(out/'review.json',{'kind':'paper_first_exploration_v1','reservoirs':reviews,'ready_jobs':ready,
            'new_optimization_evaluations':0,'source_sha256':{p:w.sha(ROOT/p) for p in ['workflows/paper_exploration.py','tools/paper_exploration.py']},
            'prior_registries_changed':False,'canonical_inputs_changed':False})
    return out
